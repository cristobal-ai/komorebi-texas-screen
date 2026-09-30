import datetime as dt
import math

import geopandas as gpd
import pandas as pd
import pytest
from shapely.geometry import box

from pipeline.common import load_config

M2_PER_ACRE = 4046.8564224


@pytest.fixture
def cfg():
    return load_config()


def gen_row(pid, name, mw, year, month=6, dc=None, ba="ERCO", county="Pecos", state="TX", src="SUN", pm="PV", gid="PV1"):
    return {
        "Entity Name": "Op " + name, "Plant ID": pid, "Plant Name": name, "Plant State": state, "County": county,
        "Balancing Authority Code": ba, "Generator ID": gid, "Nameplate Capacity (MW)": mw, "DC Net Capacity (MW)": dc,
        "Operating Month": month, "Operating Year": year, "Energy Source Code": src, "Prime Mover Code": pm,
        "Latitude": 31.0, "Longitude": -103.0, "source_month": "2026-08",
    }


@pytest.fixture
def generators():
    return pd.DataFrame([
        gen_row(1, "Alpha", 50, 2017, 6, dc=65, gid="PV1"),     # 2 phases, 80 MW, DC 104 → ILR 1.30
        gen_row(1, "Alpha", 30, 2018, 3, dc=39, gid="PV2"),
        gen_row(1, "Alpha", 20, 2019, 1, src="MWH", pm="BA", gid="BESS"),  # storage, excluded from MW
        gen_row(2, "Bravo", 150, 2016, 9),                      # no EIA DC → USPVDB DC
        gen_row(3, "Charlie", 8, 2018),                         # < 10 MW
        gen_row(4, "Delta", 60, 2023, 2),                       # COD too late
        gen_row(5, "Echo", 30, 2019),                           # 100 acres → 3.3 ac/MW
        gen_row(6, "Foxtrot", 40, 2020),                        # no polygon → review
        gen_row(7, "Golf", 20, 2021, ba="EPE", county="El Paso"),  # non-ERCOT flag, still passes
        gen_row(8, "Hotel", 100, 2015, 1, dc=130),              # boundary: exactly 100 MW, planned load 74.75
        gen_row(9, "India", 200, 2019, state="NM", county="Lea"),  # not TX
    ])


def _square(acres, x0):
    side = math.sqrt(acres * M2_PER_ACRE)
    return box(x0, 800000, x0 + side, 800000 + side)


@pytest.fixture
def polygons():
    rows = [  # eia_id, acres, ac, dc, year, county
        (1, 640, 80, 104, 2017, "Pecos County"),
        (2, 1200, 150, 195, 2016, "Pecos County"),
        (3, 70, 8, 10, 2018, "Pecos County"),
        (4, 480, 60, 78, 2023, "Pecos County"),
        (5, 100, 30, 39, 2019, "Pecos County"),
        (7, 160, 20, 26, 2021, "El Paso County"),
        (8, 800, 100, 130, 2015, "Pecos County"),
        (99, 300, 40, 52, 2020, "Pecos County"),   # orphan: not in EIA-860M
        (None, 90, 12, 15, 2018, "Pecos County"),  # no eia_id
        (1, 0, 0, 0, 2017, "Lea County"),          # different state row ignored below
    ]
    geoms, recs = [], []
    for i, (eid, acres, ac, dc, yr, cty) in enumerate(rows):
        state = "NM" if cty == "Lea County" else "TX"
        geoms.append(_square(max(acres, 1), -600000 + i * 20000))
        recs.append({"case_id": i, "eia_id": eid, "p_state": state, "p_county": cty, "p_area": acres * M2_PER_ACRE,
                     "p_cap_ac": ac, "p_cap_dc": dc, "p_year": yr, "p_axis": "single-axis", "p_tech_pri": "crystalline"})
    return gpd.GeoDataFrame(recs, geometry=geoms, crs="EPSG:5070").to_crs("EPSG:4326")


@pytest.fixture
def as_of():
    return dt.date(2026, 9, 30)


# ---- EIA-860 / EIA-923 fixtures laid out like the real files (title rows, line breaks in headers, zipped) -------
import zipfile

import openpyxl


def _sheet(wb, title, header, rows, title_rows=1):
    ws = wb.create_sheet(title)
    for i in range(title_rows):
        ws.append([f"U.S. Energy Information Administration — note row {i + 1}"])
    ws.append(header)
    for r in rows:
        ws.append(r)


def _zip(path, members):
    with zipfile.ZipFile(path, "w") as z:
        for name, wb in members.items():
            tmp = path.parent / name
            wb.save(tmp)
            z.write(tmp, name)
            tmp.unlink()
    return path


@pytest.fixture
def eia860_zip(tmp_path):
    plant = openpyxl.Workbook(); plant.remove(plant.active)
    _sheet(plant, "Plant", ["Utility ID", "Plant Code", "Plant Name", "State", "Grid Voltage (kV)",
                            "Grid Voltage 2 (kV)", "Grid Voltage 3 (kV)"], [
        [1, 1, "Alpha", "TX", 345, None, None],
        [1, 2, "Bravo", "TX", 138, 345, None],
        [1, 5, "Echo", "TX", 34.5, None, None],
        [1, 8, "Hotel", "TX", None, None, None],
    ])
    solar = openpyxl.Workbook(); solar.remove(solar.active)
    hdr = ["Plant Code", "Generator ID", "Nameplate\nCapacity (MW)", "Single-Axis Tracking?", "Dual-Axis Tracking?",
           "Fixed Tilt?", "East West Fixed Tilt?", "Bifacial?", "Crystalline Silicon?", "Thin-Film (CdTe)?"]
    _sheet(solar, "Operable", hdr, [
        [1, "PV1", 50, "Y", "N", "N", "N", "N", "Y", "N"],
        [1, "PV2", 30, "Y", "N", "N", "N", "Y", "Y", "N"],
        [2, "PV1", 150, "N", "N", "Y", "N", None, "N", "Y"],
        [5, "PV1", 20, "Y", "N", "Y", "N", "N", "Y", "N"],   # two tracking flags → mixed
        [5, "PV2", 10, "Y", "N", "N", "N", "N", "Y", "N"],
    ])
    _sheet(solar, "Retired and Canceled", hdr, [[1, "OLD", 5, "N", "N", "Y", "N", "N", "Y", "N"]])
    return _zip(tmp_path / "eia8602025.zip", {"2___Plant_Y2025.xlsx": plant, "3_3_Solar_Y2025.xlsx": solar})


def _netgen_row(pid, pm, freq, total, months=12):
    per = total / months if total is not None else None
    monthly = [per if i < months else None for i in range(12)]
    return [pid, f"P{pid}", pm, "SUN", *monthly, total, freq, 2025]


@pytest.fixture
def eia923_zip(tmp_path):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    months = [f"Netgen\n{m}" for m in ("January", "February", "March", "April", "May", "June", "July", "August",
                                         "September", "October", "November", "December")]
    hdr = ["Plant Id", "Plant Name", "Reported\nPrime Mover", "Reported\nFuel Type Code", *months,
           "Net Generation\n(Megawatthours)", "Respondent\nFrequency", "YEAR"]
    _sheet(wb, "Page 1 Generation and Fuel Data", hdr, [
        _netgen_row(1, "PV", "M", 0.28 * 80 * 8760),        # 196,224 MWh → CF 0.280, monthly
        _netgen_row(1, "BA", "M", -500.0),                  # storage row: excluded
        _netgen_row(2, "PV", "A", 0.22 * 150 * 8760),       # annual respondent
        _netgen_row(5, "PV", "M", 60000.0, months=9),       # preliminary: 9 months → incomplete
        _netgen_row(8, "PV", "M", 0.45 * 100 * 8760),       # implausible → noted, value kept
        _netgen_row(99999, "PV", "A", 1234.0),              # state-level increment row: no plant
    ], title_rows=5)
    _sheet(wb, "Page 4 Generator Data", ["Plant Id", "Generator Id"], [[1, "PV1"]], title_rows=5)
    return _zip(tmp_path / "f923_2025.zip", {"EIA923_Schedules_2_3_4_5_M_12_2025_Final.xlsx": wb})
